import base64
import html
from io import BytesIO
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "art-assets-catalog.docx"
HTML_OUTPUT = ROOT / "docs" / "art-assets-catalog.html"
PUBLIC_BASE = "https://fafa-big-faker.github.io/growth-partner/"


def asset_specs():
    return [
        (
            "背景",
            ROOT / "assets/images/v2/backgrounds",
            ["login.webp", "cultivate.webp", "tasks.webp", "reward.webp"],
            lambda _name: "建议 1400×1038 或更高，近 4:3 横图，无文字人物；最终导出 700×519 WebP",
        ),
        (
            "角色待机动画",
            ROOT / "assets/images/character/idle",
            [f"frame-{index:02d}.png" for index in range(1, 7)],
            lambda _name: "严格 362×724，透明背景，角色基线和中心一致，不裁切；同组共 6 帧",
        ),
        (
            "角色砍树动画",
            ROOT / "assets/images/character/chop",
            [f"frame-{index:02d}.png" for index in range(1, 7)],
            lambda _name: "严格 362×724，透明背景，角色基线和中心一致，不裁切；同组共 6 帧",
        ),
        (
            "仙树",
            ROOT / "assets/images/v2/trees",
            ["sprout.png", "spirit.png", "divine.png"],
            lambda _name: "透明背景；建议统一 512×640 画布，脚底基线一致，主体完整",
        ),
        (
            "功能图标",
            ROOT / "assets/images/v2/icons",
            [
                "icon-achievement.png", "icon-breakthrough.png", "icon-close.png",
                "icon-cultivate.png", "icon-forge.png", "icon-lock.png", "icon-mail.png",
                "icon-reward.png", "icon-shop.png", "icon-tasks.png", "icon-tree-info.png",
                "icon-wallet.png",
            ],
            lambda _name: "透明背景，建议 256×256 正方形画布，主体居中，适合 24–64px 显示",
        ),
        (
            "特效",
            ROOT / "assets/images/v2/effects",
            ["effect-drop-glow.png", "effect-hit-spark.png", "effect-leaf-gold.png", "effect-leaf-green.png"],
            lambda _name: "透明背景，建议 256×256 正方形画布，边缘保留透明空间，无硬底色",
        ),
        (
            "界面组件",
            ROOT / "assets/images/v2/ui",
            [
                "button-primary.png", "button-secondary.png", "checkbox-off.png", "checkbox-on.png",
                "modal-crest.png", "panel-corner.png", "panel-divider.png", "scroll-thumb.png",
                "slot-blue.png", "slot-gold.png", "slot-neutral.png", "slot-purple.png", "slot-rose.png",
                "status-pill.png", "tab-active.png", "tab-inactive.png",
            ],
            ui_requirement,
        ),
        (
            "道具图标",
            ROOT / "assets/images/icons",
            [
                "0.png", "1.png", "10001.png", "10002.png", "10101.png", "10102.png",
                "10201.png", "10202.png", "10301.png", "10302.png", "20001.png", "20101.png",
                "20201.png", "20301.png", "30001.png", "30101.png", "30201.png", "40001.png",
                "40002.png", "51001.png", "51002.png", "52001.png", "52002.png", "53001.png",
                "53002.png", "54001.png", "54002.png", "55001.png",
            ],
            item_requirement,
        ),
    ]


def ui_requirement(name):
    if name.startswith("button-"):
        return "建议 320×160，透明背景，中央区域可九宫格拉伸"
    if name.startswith("checkbox-"):
        return "建议 128×128，透明背景，两种状态轮廓完全一致"
    if name.startswith("slot-"):
        return "建议 256×256，透明背景，边框完整且不越界"
    if name.startswith("tab-"):
        return "建议 320×160，透明背景，两种状态轮廓一致"
    if name == "scroll-thumb.png":
        return "建议 96×256，透明背景，适合纵向拉伸"
    if name == "status-pill.png":
        return "建议 192×128，透明背景，中央区域可横向拉伸"
    return "保持当前宽高比例，透明背景，装饰边缘保留安全空间"


def item_requirement(name):
    item_id = int(Path(name).stem)
    if 51000 <= item_id <= 55999:
        return "严格 64×94，透明背景，斧柄垂直，主体完整并留 2–4px 安全边距"
    return "严格 64×64，透明背景，主体居中并留 2–4px 安全边距"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    table_header = OxmlElement("w:tblHeader")
    table_header.set(qn("w:val"), "true")
    tr_pr.append(table_header)


def set_font(run, name="Microsoft YaHei", size=9, bold=False, color="1F2933"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_thumbnail(cell, image_path):
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    with Image.open(image_path) as image:
        width_px, height_px = image.size
        converted = image.convert("RGBA" if "A" in image.getbands() else "RGB")
        image_stream = BytesIO()
        converted.save(image_stream, format="PNG")
        image_stream.seek(0)
    max_width = 0.82
    max_height = 0.78
    scale = min(max_width / width_px, max_height / height_px)
    width = max(0.16, width_px * scale)
    height = max(0.16, height_px * scale)
    run = paragraph.add_run()
    run.add_picture(image_stream, width=Inches(width), height=Inches(height))


def add_text(cell, text, align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, color="1F2933"):
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(text)
    set_font(run, size=8.5, bold=bold, color=color)


def style_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.05)
    table.columns[1].width = Inches(1.75)
    table.columns[2].width = Inches(4.35)
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.width = (Inches(1.05), Inches(1.75), Inches(4.35))[col_index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, "355C68")
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F3F7F6")


def build_document():
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(9)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(7)
    run = title.add_run("寻道大千美术资源清单")
    set_font(run, size=24, bold=True, color="000000")

    intro = document.add_paragraph()
    intro.paragraph_format.space_after = Pt(14)
    intro_run = intro.add_run(
        "本清单收录当前游戏已上线使用的 79 个美术文件，并给出替换资源的推荐尺寸与制作约束。"
        "替换时请保持文件名不变；透明背景、主体安全边距和同组动画基线需要严格一致。"
    )
    set_font(intro_run, size=10, color="374151")

    groups = asset_specs()
    total = 0
    for group_index, (title_text, directory, filenames, requirement) in enumerate(groups):
        if group_index > 0:
            document.add_page_break()
        heading = document.add_paragraph(style="Heading 1")
        heading.paragraph_format.space_before = Pt(0)
        heading.paragraph_format.space_after = Pt(8)
        heading_run = heading.add_run(f"{title_text}  {len(filenames)} 项")
        set_font(heading_run, size=16, bold=True, color="000000")

        table = document.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        headers = ("图像", "图像名", "尺寸要求")
        for index, header in enumerate(headers):
            add_text(table.rows[0].cells[index], header, WD_ALIGN_PARAGRAPH.CENTER, True, "FFFFFF")
        set_repeat_table_header(table.rows[0])

        for filename in filenames:
            image_path = directory / filename
            if not image_path.exists():
                raise FileNotFoundError(image_path)
            row = table.add_row()
            add_thumbnail(row.cells[0], image_path)
            add_text(row.cells[1], filename, WD_ALIGN_PARAGRAPH.CENTER, True)
            add_text(row.cells[2], requirement(filename))
            total += 1
        style_table(table)

    if total != 79:
        raise ValueError(f"Expected 79 assets, found {total}")

    document.core_properties.title = "寻道大千美术资源清单"
    document.core_properties.subject = "当前上线美术资源及替换尺寸要求"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


def build_html():
    sections = []
    for title_text, directory, filenames, requirement in asset_specs():
        rows = []
        for filename in filenames:
            image_path = directory / filename
            with Image.open(image_path) as source:
                thumbnail = source.convert("RGBA" if "A" in source.getbands() else "RGB")
                thumbnail.thumbnail((180, 180), Image.Resampling.LANCZOS)
                image_stream = BytesIO()
                thumbnail.save(image_stream, format="PNG", optimize=True)
            embedded_image = base64.b64encode(image_stream.getvalue()).decode("ascii")
            rows.append(
                "<tr>"
                f'<td class="preview"><img src="data:image/png;base64,{embedded_image}" alt="{html.escape(filename)}"></td>'
                f'<td class="name">{html.escape(filename)}</td>'
                f'<td>{html.escape(requirement(filename))}</td>'
                "</tr>"
            )
        sections.append(
            f"<h2>{html.escape(title_text)} <span>{len(filenames)} 项</span></h2>"
            "<table><thead><tr><th>图像</th><th>图像名</th><th>尺寸要求</th></tr></thead>"
            f"<tbody>{''.join(rows)}</tbody></table>"
        )
    page = f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><title>寻道大千美术资源清单</title>
<style>
body{{font-family:'Microsoft YaHei',sans-serif;color:#1f2933;max-width:960px;margin:36px auto;padding:0 28px;background:#fff}}
h1{{font-size:30px;margin:0 0 10px;color:#111}}p{{color:#4b5563;line-height:1.75;margin:0 0 28px}}
h2{{font-size:20px;color:#111;margin:30px 0 10px}}h2 span{{font-size:13px;color:#607d78;font-weight:500}}
table{{width:100%;border-collapse:collapse;table-layout:fixed;margin-bottom:30px}}
th,td{{border:1px solid #d9d9d9;padding:10px 12px;vertical-align:middle;line-height:1.55}}
th{{background:#355c68;color:#fff;text-align:center}}tbody tr:nth-child(even){{background:#f3f7f6}}
th:nth-child(1),td:nth-child(1){{width:108px}}th:nth-child(2),td:nth-child(2){{width:190px}}
.preview{{text-align:center;height:92px}}.preview img{{max-width:86px;max-height:86px;object-fit:contain}}
.name{{text-align:center;font-weight:600;word-break:break-all}}
</style></head><body><h1>寻道大千美术资源清单</h1>
<p>本清单收录当前游戏已上线使用的 79 个美术文件，并给出替换资源的推荐尺寸与制作约束。替换时请保持文件名不变；透明背景、主体安全边距和同组动画基线需要严格一致。</p>
{''.join(sections)}</body></html>"""
    HTML_OUTPUT.write_text(page, encoding="utf-8")
    print(HTML_OUTPUT)


if __name__ == "__main__":
    build_document()
    build_html()
